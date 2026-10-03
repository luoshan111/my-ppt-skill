# -*- coding: utf-8 -*-
"""模式 B：模板克隆流水线。基于原 .pptx 在 XML 层克隆页面并预算内替换文字，
视觉与原版 100% 一致（渐变/自由曲线/嵌入字体/母版全部保留）。

用法:
    python clone_deck.py <template.pptx> <out.pptx> <plan.json>

plan.json 结构（pages 顺序即输出顺序；未引用的模板页全部删除）:
{
  "pages": [
    {"src": 1, "texts": {"原文精确段落": "新文字", "整框原文\\n第二行": "新文字\\n第二行"}},
    {"src": 4},
    {"src": 4, "texts": {...}}          // 同一源页可克隆多次
  ]
}

规则与限制:
- 文字替换预算 = 原文长度 × 1.15，超预算仅告警不阻断（渲染后人工确认）
- 支持普通文本框与表格单元格；图表内文字不支持
- 原文件如嵌入字体子集，替换后的新字符可能缺字形回退系统字体（见 SKILL.md 模式 B）
"""
import copy, json, sys
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

BUDGET = 1.15

def _norm(s):
    return s.replace("\x0b", "\n").replace("\r", "").strip()

def iter_text_frames(shapes):
    for sh in shapes:
        if sh.shape_type == 6:  # GROUP
            yield from iter_text_frames(sh.shapes)
        elif getattr(sh, "has_table", False) and sh.has_table:
            for cell in sh.table.iter_cells():
                yield sh, cell.text_frame
        elif sh.has_text_frame:
            yield sh, sh.text_frame

def set_para(p, new_text):
    runs = p.runs
    if not runs:
        r = p.add_run()
        r.text = new_text
        return
    runs[0].text = new_text
    for r in runs[1:]:
        r._r.getparent().remove(r._r)

def apply_texts(slide, mapping):
    """整框匹配优先，其次段级匹配。返回 (hits, warns)。"""
    m = {_norm(k): v for k, v in mapping.items()}
    hits, warns = [], []
    for sh, tf in iter_text_frames(slide.shapes):
        paras = list(tf.paragraphs)
        full = "\n".join(p.text for p in paras)
        if _norm(full) in m:
            new = m[_norm(full)]
            if len(new) > len(full) * BUDGET:
                warns.append(f"超预算: '{full[:12]}...' {len(full)}->{len(new)}")
            parts = new.split("\n")
            if len(parts) > len(paras):
                # 新文本行数多于段落数：尾部行以软换行(\x0b)并入最后一段
                for i, p in enumerate(paras[:-1]):
                    set_para(p, parts[i])
                set_para(paras[-1], "\x0b".join(parts[len(paras) - 1:]))
            else:
                for i, p in enumerate(paras):
                    set_para(p, parts[i] if i < len(parts) else "")
            hits.append(_norm(full)[:14])
            continue
        for p in paras:
            old = _norm(p.text)
            if old in m:
                new = m[old]
                if len(new) > max(len(p.text), 1) * BUDGET and len(p.text) >= 6:
                    warns.append(f"超预算: '{p.text[:12]}' {len(p.text)}->{len(new)}")
                set_para(p, new)
                hits.append(old[:14])
    return hits, warns

def duplicate_slide(prs, index):
    """同包内克隆页：复制形状树 + 逐元素重挂 rel（图片/媒体保持原部件）。"""
    src = prs.slides[index]
    dest = prs.slides.add_slide(src.slide_layout)
    for sh in list(dest.shapes):
        sh._element.getparent().remove(sh._element)
    rmap = {}
    for rel in list(src.part.rels.values()):
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        if rel.is_external:
            rmap[rel.rId] = dest.part.rels.get_or_add_ext_rel(rel.reltype, rel.target_ref)
        else:
            rmap[rel.rId] = dest.part.relate_to(rel.target_part, rel.reltype)
    for child in src.shapes._spTree:
        if child.tag in (qn("p:nvGrpSpPr"), qn("p:grpSpPr")):
            continue
        el = copy.deepcopy(child)
        for e in el.iter():
            for attr in (qn("r:embed"), qn("r:link"), qn("r:id")):
                v = e.get(attr)
                if v and v in rmap:
                    e.set(attr, rmap[v])
        dest.shapes._spTree.append(el)
    return dest

def delete_slide(prs, idx):
    sld = list(prs.slides._sldIdLst)[idx]
    prs.part.drop_rel(sld.get(qn("r:id")))
    prs.slides._sldIdLst.remove(sld)

def strip_embedded_fonts(prs):
    """移除演示文稿的嵌入字体清单：全文统一回退到本机已装字体，
    避免「原字符用嵌入子集、新字符回退系统字体」的混排异样。"""
    el = prs.part._element
    removed = 0
    for lst in el.findall(qn("p:embeddedFontLst")):
        el.remove(lst)
        removed += 1
    return removed

def main(template, out, plan_path):
    strip = "--strip-fonts" in sys.argv
    plan = json.load(open(plan_path, encoding="utf-8"))
    prs = Presentation(template)
    if strip:
        n = strip_embedded_fonts(prs)
        print(f"已移除嵌入字体清单 {n} 处（全文统一回退本机字体）")
    n_orig = len(prs.slides._sldIdLst)
    produced = []
    all_warns = []
    for i, spec in enumerate(plan["pages"]):
        src = spec["src"] - 1
        if not 0 <= src < n_orig:
            sys.exit(f"plan 错误: src={spec['src']} 超出模板页数 {n_orig}")
        duplicate_slide(prs, src)
        new_idx = n_orig + i
        slide = prs.slides[new_idx]
        hits, warns = apply_texts(slide, spec.get("texts", {}))
        all_warns += [f"page{i+1}(src{spec['src']}) {w}" for w in warns]
        missing = [k for k in spec.get("texts", {}) if not any(k[:14] in h for h in hits)]
        produced.append(f"page{i+1}<=src{spec['src']} 替换{len(hits)}处 未命中{len(missing)}")
        for k in missing:
            all_warns.append(f"page{i+1} 未命中替换键: '{k[:30]}...'")
    for idx in range(n_orig - 1, -1, -1):   # 删除全部原始页（高索引先删）
        delete_slide(prs, idx)
    prs.save(out)
    print(f"克隆完成: {len(plan['pages'])} 页 -> {out}")
    print("\n".join(produced))
    if all_warns:
        print("⚠ 告警:")
        print("\n".join("  " + w for w in all_warns))

if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2], sys.argv[3])
