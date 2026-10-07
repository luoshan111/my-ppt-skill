# -*- coding: utf-8 -*-
"""程序化 QA：文本溢出估计 + 越界 + 文本框重叠 + 信息密度。用法：
    python check.py "输出文件.pptx" [--min-units 8]
  --min-units N：每页信息单元（含字形的形状+文本框）少于 N 时报 [密度] 警告（内容页建议 8，竞赛图解页 15-30）
"""
import math, sys
from pptx import Presentation

args = [a for a in sys.argv[1:] if not a.startswith("--")]
MIN_UNITS = int(sys.argv[sys.argv.index("--min-units") + 1]) if "--min-units" in sys.argv else 0
PATH = args[0] if args else "out.pptx"
EMU_IN = 914400.0

def char_w(ch, pt):
    if ord(ch) > 0x2E7F: return pt * 1.02
    if ch in "iIl.,:;!|'\"()[] ": return pt * 0.30
    return pt * 0.55

def para_lines(text, pt, box_w):
    if not text.strip(): return 1
    w = sum(char_w(c, pt) for c in text) / 72.0
    return max(1, math.ceil(w / max(box_w, 0.1) - 0.02))

prs = Presentation(PATH)
SW, SH = prs.slide_width / EMU_IN, prs.slide_height / EMU_IN
issues = []
for si, slide in enumerate(prs.slides, 1):
    boxes = []
    for sh in slide.shapes:
        if not sh.has_text_frame: continue
        tf = sh.text_frame
        txt = "\n".join(p.text for p in tf.paragraphs)
        if not txt.strip(): continue
        x, y = sh.left / EMU_IN, sh.top / EMU_IN
        w, h = sh.width / EMU_IN, sh.height / EMU_IN
        if x < -0.02 or y < -0.02 or x + w > SW + 0.02 or y + h > SH + 0.02:
            issues.append(f"[越界] slide{si} '{txt[:14]}' ({x:.2f},{y:.2f},{w:.2f},{h:.2f})")
        total = 0.0
        for p in tf.paragraphs:
            pts = [r.font.size.pt for r in p.runs if r.font.size] or [12]
            total += para_lines(p.text, max(pts), w) * max(pts) * 1.35 / 72.0
        if total > h * 1.12 + 0.06:
            issues.append(f"[溢出?] slide{si} '{txt[:14]}' 需{total:.2f}in 框{h:.2f}in")
        boxes.append((x, y, w, h, txt))
    for i in range(len(boxes)):
        for j in range(i + 1, len(boxes)):
            a, b = boxes[i], boxes[j]
            ix = max(0, min(a[0]+a[2], b[0]+b[2]) - max(a[0], b[0]))
            iy = max(0, min(a[1]+a[3], b[1]+b[3]) - max(a[1], b[1]))
            if ix * iy > 0.05:
                issues.append(f"[重叠] slide{si} '{a[4][:10]}' x '{b[4][:10]}'")
print(f"slides={len(prs.slides._sldIdLst)}")
if MIN_UNITS:
    # 信息单元 = 有字的文本框 + 图片 + 表格/图表 + 组合；纯装饰色块（渐变切片/底带）不计
    for si, slide in enumerate(prs.slides, 1):
        units = 0
        for sh in slide.shapes:
            if sh.shape_type is not None and str(sh.shape_type).startswith("GROUP"): units += 1
            elif sh.has_text_frame and sh.text_frame.text.strip(): units += 1
            elif sh.shape_type is not None and ("PICTURE" in str(sh.shape_type) or "TABLE" in str(sh.shape_type) or "CHART" in str(sh.shape_type)): units += 1
        if units < MIN_UNITS:
            issues.append(f"[密度] slide{si} 仅 {units} 个信息单元（要求 ≥{MIN_UNITS}）——内容页应紧凑丰富，用 fillStack/gridAreas 撑满区域")
print("\n".join(issues) if issues else "NO ISSUES")
